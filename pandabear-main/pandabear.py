import openpyxl
import pandas as pd



local_var = "Bruce"
result_dict = {}

# 1. Load the Excel file with openpyxl to inspect cell colors
wb = openpyxl.load_workbook("C:/Users/bruce/Downloads/Book 1.xlsx", data_only=True)
sheet = wb.active

# 1. Find the row matching your local_var in column 1
target_row_idx = None
for row in range(2, sheet.max_row + 1):
    if sheet.cell(row=row, column=1).value == local_var:
        target_row_idx = row
        break

# 2. If found, process the timeline based on theme index 9
if target_row_idx:
    hours_timeline = []
    
    for col in range(2, sheet.max_column + 1):
        # Fallback if your header row happens to be empty
        hour_header = sheet.cell(row=1, column=col).value
        if hour_header is None:
            hour_header = f"Hour {col - 1}"
            
        cell = sheet.cell(row=target_row_idx, column=col)
        color_obj = cell.fill.start_color
        
        # Check if it uses the green theme index (9)
        if color_obj.type == "theme" and color_obj.theme == 9:
            status = "Occupied"
        else:
            status = "Available"
            
        hours_timeline.append(f"{hour_header}: {status}")
        
    result_dict[local_var] = hours_timeline

print("Final Result:")
print(result_dict)


#check against json db


#1 grab all members into a list
#2 for all members in list, iterate through via indexing of their row to then dynamically assign values to the member inside the dictionary_conversion
#3 once dictionary is is finished -> export utilizing proper formatting
# Note -> the dictionary is already the formatting of the "Shifts" tab of the excel sheet, and while other tabs must be included in the export, 
#     they can be left blank 
#   Also need to grab the email based on the members name that has been configged by the supervisor already which stores the connections in a local data structure
#   Need a function to determine days found based on white space arithmetic, and a function that runs after a day has been found due to holidays and dead days
#   My approach will be to loop based on start and stop coord so theres no excessive function guessing

# member: str
# work_email: str
# group = "CSR" 
# start_date: str 
# start_time:str 
# end_date: str 
# end_time: str
# theme_color: str 
# cutom_label = None
# unpaid_break = None 
# notes = None
# shared = None

# dictionary_conversion = {
#   member: [work_email, group, start_date, start_time, end_date, end_time, theme_color, cutom_label, unpaid_break, notes, shared]
# }


# file explorer opens and allows user to navigate to find the sheet which they wish to convert
#1 get amount of rows and columns from list based on fixed input as a starting
#2 double iterate through each name/row and every coloumn on the row before moving to the next one while storing 
#   pertinent info in accurate temp dictionary
#3 convert dictionary into adaquate excel sheet format for importation into the downloads folder


#notes - use regex on last name to connect email. Store a small database of 
# use live calendar to determine which sheets to convert based off the massive main excel sheet
# access to a local dictionary database of emails and first and last name values for easy referencing

# Loop through rows (skip header if row 1 is headers)